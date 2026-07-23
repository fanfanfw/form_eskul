from fastapi import FastAPI, Request, Form, HTTPException, UploadFile, File, Body
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from fastapi.staticfiles import StaticFiles
from starlette.middleware.sessions import SessionMiddleware
import psycopg2
from psycopg2.extras import RealDictCursor
import os
import io
import csv
import hmac
import re
import hashlib
import zipfile
from collections import Counter
from typing import List
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

try:
    from config import *
except ImportError:
    DB_HOST = "localhost"
    DB_DATABASE = "db_form_eskul"
    DB_USER = "postgres"
    DB_PASSWORD = "postgres"
    DB_PORT = 5432
    APP_TITLE = "Form Eskul Siswa"
    APP_DESCRIPTION = "Website untuk mengisi form eskul siswa"

DB_HOST = os.getenv("DB_HOST", DB_HOST)
DB_DATABASE = os.getenv("DB_DATABASE", DB_DATABASE)
DB_USER = os.getenv("DB_USER", DB_USER)
DB_PASSWORD = os.getenv("DB_PASSWORD", DB_PASSWORD)
DB_PORT = int(os.getenv("DB_PORT", DB_PORT))
DATABASE_URL = os.getenv("DATABASE_URL")
ADMIN_PIN = os.getenv("ADMIN_PIN")
SESSION_SECRET = os.getenv("SESSION_SECRET") or os.urandom(32).hex()

app = FastAPI(title=APP_TITLE, description=APP_DESCRIPTION)
templates = Jinja2Templates(directory="templates")
app.mount("/static", StaticFiles(directory="static"), name="static")

DB_CONFIG = {"host": DB_HOST, "database": DB_DATABASE, "user": DB_USER, "password": DB_PASSWORD, "port": DB_PORT}
ADMIN_PATHS = (
    "/api/students/preview-import", "/api/students/import", "/api/students/manage",
    "/api/students/create", "/api/students/bulk-delete", "/api/eskul/manage",
    "/api/eskul/create", "/api/registrations", "/api/registrations/export"
)

@app.middleware("http")
async def protect_admin(request: Request, call_next):
    path = request.url.path
    protected = path == "/registrations" or any(path == prefix or path.startswith(prefix + "/") for prefix in ADMIN_PATHS) or (path.startswith("/api/students/") and path.endswith(("/update",))) or (path.startswith("/api/students/") and request.method == "DELETE") or (path.startswith("/api/eskul/") and path != "/api/eskul")
    if protected and not request.session.get("admin"):
        if path == "/registrations":
            return templates.TemplateResponse("registrations.html", {"request": request, "login": True, "admin_configured": bool(ADMIN_PIN)}, status_code=401)
        return HTMLResponse("Admin authentication required", status_code=401)
    return await call_next(request)

app.add_middleware(SessionMiddleware, secret_key=SESSION_SECRET, same_site="lax", https_only=os.getenv("SESSION_HTTPS_ONLY", "true").lower() == "true")

def get_db_connection():
    try:
        if DATABASE_URL:
            return psycopg2.connect(DATABASE_URL, cursor_factory=RealDictCursor)
        return psycopg2.connect(**DB_CONFIG, cursor_factory=RealDictCursor)
    except psycopg2.Error:
        return None

def db_or_500():
    conn = get_db_connection()
    if not conn:
        raise HTTPException(status_code=500, detail="Database connection failed")
    return conn

@app.on_event("startup")
def migrate_database():
    conn = db_or_500()
    try:
        cursor = conn.cursor()
        cursor.execute("SELECT pg_advisory_xact_lock(hashtext('form_eskul_schema_migration'))")
        cursor.execute("SELECT 1 FROM pg_attribute WHERE attrelid='eskul'::regclass AND attname='minimal_kelas' AND NOT attisdropped")
        new_column = not cursor.fetchone()
        cursor.execute("ALTER TABLE eskul ADD COLUMN IF NOT EXISTS minimal_kelas INTEGER NOT NULL DEFAULT 1")
        if new_column:
            cursor.execute("UPDATE eskul SET minimal_kelas=3 WHERE LOWER(nama_eskul) IN ('pencak silat','futsal','angklung')")
        cursor.execute("SELECT 1 FROM pg_constraint WHERE conname='eskul_minimal_kelas_check' AND conrelid='eskul'::regclass")
        if not cursor.fetchone():
            cursor.execute("ALTER TABLE eskul ADD CONSTRAINT eskul_minimal_kelas_check CHECK (minimal_kelas BETWEEN 1 AND 6)")
        conn.commit()
    finally:
        conn.close()

def grade_from_kelas(kelas):
    match = re.search(r"\d+", kelas or "")
    if not match or not 1 <= int(match.group()) <= 6:
        raise HTTPException(400, "Kelas tidak valid")
    return int(match.group())

def validate_assignment(cursor, kelas, eskul_id):
    if eskul_id is None:
        return
    grade = grade_from_kelas(kelas)
    cursor.execute("SELECT minimal_kelas FROM eskul WHERE id=%s FOR SHARE", (eskul_id,))
    eskul = cursor.fetchone()
    if not eskul:
        raise HTTPException(400, "Eskul tidak valid")
    if grade < eskul["minimal_kelas"]:
        raise HTTPException(400, f"Eskul hanya untuk kelas {eskul['minimal_kelas']}+")

def clean_student(nis, nisn, nama, jeniskelamin, kelas):
    values = [str(value).strip() for value in (nis, nisn, nama, jeniskelamin, kelas)]
    values[3] = values[3].upper()
    if not values[0] or not values[2] or not values[4]:
        raise HTTPException(status_code=400, detail="NIS, nama, dan kelas wajib diisi")
    if values[3] not in ("L", "P", "-"):
        raise HTTPException(status_code=400, detail="Jenis kelamin harus L, P, atau -")
    limits = (20, 20, 255, 1, 15)
    labels = ("NIS", "NISN", "Nama", "Jenis kelamin", "Kelas")
    for value, limit, label in zip(values, limits, labels):
        if len(value) > limit:
            raise HTTPException(status_code=400, detail=f"{label} maksimal {limit} karakter")
    return values

def filters_sql(search="", kelas="", eskul_id=None, status=""):
    clauses, params = [], []
    search = search.strip()
    kelas = kelas.strip()
    if search:
        clauses.append("(s.nis ILIKE %s OR s.nisn ILIKE %s OR s.nama ILIKE %s)")
        params.extend([f"%{search}%"] * 3)
    if kelas:
        clauses.append("s.kelas = %s")
        params.append(kelas)
    if eskul_id is not None:
        clauses.append("s.eskul = %s")
        params.append(eskul_id)
    if status == "assigned":
        clauses.append("s.eskul IS NOT NULL")
    elif status == "unassigned":
        clauses.append("s.eskul IS NULL")
    elif status:
        raise HTTPException(status_code=400, detail="Status tidak valid")
    return (" WHERE " + " AND ".join(clauses) if clauses else ""), params

@app.get("/", response_class=HTMLResponse)
async def home(request: Request):
    return templates.TemplateResponse("index.html", {"request": request})

@app.post("/registrations/login")
async def admin_login(request: Request, pin: str = Form(...)):
    if not ADMIN_PIN or not hmac.compare_digest(pin.encode(), ADMIN_PIN.encode()):
        return templates.TemplateResponse("registrations.html", {"request": request, "login": True, "admin_configured": bool(ADMIN_PIN), "login_error": "PIN salah atau akses admin belum dikonfigurasi"}, status_code=401)
    request.session["admin"] = True
    return RedirectResponse("/registrations", status_code=303)

@app.post("/registrations/logout")
async def admin_logout(request: Request):
    request.session.clear()
    return RedirectResponse("/registrations", status_code=303)

@app.get("/api/kelas")
async def get_kelas():
    conn = db_or_500()
    try:
        cursor = conn.cursor(); cursor.execute("SELECT DISTINCT kelas FROM siswa ORDER BY kelas")
        return {"kelas": [row["kelas"] for row in cursor.fetchall()]}
    finally: conn.close()

@app.get("/api/siswa/{kelas}")
async def get_siswa_by_kelas(kelas: str):
    conn = db_or_500()
    try:
        cursor = conn.cursor(); cursor.execute("SELECT id, nis, nisn, nama FROM siswa WHERE kelas = %s AND eskul IS NULL ORDER BY nama", (kelas,))
        return {"siswa": cursor.fetchall()}
    finally: conn.close()

@app.get("/api/eskul")
async def get_eskul():
    conn = db_or_500()
    try:
        cursor = conn.cursor(); cursor.execute("SELECT id, nama_eskul, minimal_kelas FROM eskul ORDER BY nama_eskul")
        return {"eskul": cursor.fetchall()}
    finally: conn.close()

def excel_text(value):
    if value is None: return ""
    if isinstance(value, float) and value.is_integer(): return str(int(value))
    return str(value).strip().lstrip("'").strip()

def header_key(value):
    return re.sub(r"[^a-z0-9]", "", excel_text(value).lower())

def normalized_class(value):
    text = excel_text(value)
    match = re.fullmatch(r"\s*(?:kelas|class)\s*([1-6])\s*[- ]?([a-z])?\s*", text, re.I) or re.fullmatch(r"\s*([1-6])\s*[- ]?([a-z])?\s*", text, re.I)
    return f"Kelas {match.group(1)}{match.group(2).upper() if match.group(2) else ''}" if match else ""

def normalized_gender(value):
    key = header_key(value)
    if key in ("l", "lakilaki", "laki", "male", "pria"): return "L"
    if key in ("p", "perempuan", "female", "wanita"): return "P"
    return "-"

async def read_limited_upload(upload, maximum=15 * 1024 * 1024):
    content = bytearray()
    while chunk := await upload.read(1024 * 1024):
        content.extend(chunk)
        if len(content) > maximum: raise HTTPException(400, "File maksimal 15MB")
    return bytes(content)

def read_students_excel(file_content: bytes):
    try:
        with zipfile.ZipFile(io.BytesIO(file_content)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000: raise HTTPException(400, "File Excel memiliki terlalu banyak entri")
            if sum(entry.file_size for entry in entries) > 100 * 1024 * 1024: raise HTTPException(400, "Isi file Excel terlalu besar")
            if any(entry.file_size > 50 * 1024 * 1024 or entry.file_size and (not entry.compress_size or entry.file_size / entry.compress_size > 200) for entry in entries): raise HTTPException(400, "File Excel terkompresi mencurigakan")
    except (zipfile.BadZipFile, OSError, ValueError):
        raise HTTPException(400, "File Excel rusak, terenkripsi, atau tidak dapat dibaca")
    try: workbook = load_workbook(io.BytesIO(file_content), read_only=True, data_only=True)
    except (InvalidFileException, OSError, ValueError, KeyError, EOFError, zipfile.BadZipFile): raise HTTPException(400, "File Excel rusak, terenkripsi, atau tidak dapat dibaca")
    sheets = [sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
    if len(sheets) > 100: raise HTTPException(400, "File maksimal 100 sheet")
    aliases = {"name": {"nama", "namasiswa", "namalengkap"}, "gender": {"jk", "lp", "jeniskelamin"}, "class": {"kelas", "class"}, "nis": {"nis", "nomorinduksiswa"}, "nisn": {"nisn", "nomorinduksiswanasional"}, "combined": {"nisnnis", "nisnisn"}, "sequence": {"no", "nomor", "urut"}}
    rows, duplicates, skipped, summaries, seen, total_cells, synthetic_counts = [], [], [], [], set(), 0, Counter()
    for sheet in sheets:
        if sheet.max_row is None or sheet.max_column is None: sheet.calculate_dimension(force=True)
        max_row, max_column = sheet.max_row or 0, sheet.max_column or 0
        if max_row > 20000 or max_column > 200: raise HTTPException(400, "Sheet maksimal 20000 baris dan 200 kolom")
        total_cells += max_row * max_column
        if total_cells > 1000000: raise HTTPException(400, "File maksimal 1000000 sel")
        top = list(sheet.iter_rows(min_row=1, max_row=min(30, max_row), max_col=min(200, max_column), values_only=True))
        sheet_class = normalized_class(sheet.title)
        found, header_rows = {}, []
        for number, row in enumerate(top, 1):
            hits = {}
            for column, cell in enumerate(row):
                key = header_key(cell)
                for field, names in aliases.items():
                    if key in names: hits[field] = column
            if "name" in hits and (set(hits) - {"name"} or sheet_class and len(hits) == 1 and sum(bool(excel_text(cell)) for cell in row) == 1):
                found.update(hits); header_rows.append(number)
                for nearby_number in range(max(1, number - 2), min(len(top), number + 2) + 1):
                    for column, cell in enumerate(top[nearby_number - 1]):
                        key = header_key(cell)
                        for field, names in aliases.items():
                            if key in names and field not in found: found[field] = column; header_rows.append(nearby_number)
                break
        if "name" not in found:
            skipped.append({"sheet": sheet.title, "row": 0, "reason": "Header nama tidak ditemukan"}); summaries.append({"sheet": sheet.title, "class": sheet_class, "rows": 0}); continue
        count, blank_names = 0, 0
        start = max(header_rows) + 1
        for number, values in enumerate(sheet.iter_rows(min_row=start, max_row=min(max_row, 20000), max_col=min(max_column, 200), values_only=True), start):
            name = excel_text(values[found["name"]] if found["name"] < len(values) else "")
            if not name:
                blank_names += 1
                if "sequence" not in found and blank_names >= 20: break
                continue
            blank_names = 0
            if "sequence" in found and not excel_text(values[found["sequence"]] if found["sequence"] < len(values) else "").isdigit(): continue
            name_key = header_key(name)
            if name_key in aliases["name"] or name_key.startswith(("jumlah", "total", "mengetahui", "walikelas", "keterangan", "lakilaki", "perempuan")): continue
            row_class = excel_text(values[found["class"]]) if "class" in found and found["class"] < len(values) else sheet_class
            row_class = normalized_class(row_class)
            if not row_class:
                skipped.append({"sheet": sheet.title, "row": number, "reason": "Kelas tidak valid"}); continue
            nis = excel_text(values[found["nis"]]) if "nis" in found and found["nis"] < len(values) else ""
            nisn = excel_text(values[found["nisn"]]) if "nisn" in found and found["nisn"] < len(values) else ""
            if "combined" in found and found["combined"] < len(values):
                parts = [excel_text(part) for part in re.split(r"\s*/\s*", excel_text(values[found["combined"]]), maxsplit=1)]
                if len(parts) == 2: nisn, nis = parts
            gender = normalized_gender(values[found["gender"]] if "gender" in found and found["gender"] < len(values) else "")
            if gender == "-":
                skipped.append({"sheet": sheet.title, "row": number, "reason": "Jenis kelamin wajib L atau P"}); continue
            if not nis:
                base = f"{header_key(row_class)}|{header_key(name)}"
                synthetic_counts[base] += 1
                nis = "IMP-" + hashlib.sha256(f"{base}|{synthetic_counts[base]}".encode()).hexdigest()[:16]
            try:
                nis, nisn, name, gender, row_class = clean_student(nis, nisn, name, gender, row_class)
            except HTTPException as error:
                skipped.append({"sheet": sheet.title, "row": number, "reason": error.detail}); continue
            student = {"nis": nis, "nisn": nisn, "nama": name, "jeniskelamin": gender, "kelas": row_class}
            if nis in seen:
                duplicates.append({**student, "sheet": sheet.title, "row": number}); continue
            seen.add(nis); rows.append(student); count += 1
            if len(rows) > 10000: raise HTTPException(400, "File maksimal 10000 siswa")
        summaries.append({"sheet": sheet.title, "class": sheet_class, "rows": count})
    if not rows: raise HTTPException(400, "Tidak ada baris siswa yang dapat digunakan")
    return {"rows": rows, "duplicates": duplicates, "skipped": skipped, "sheets": summaries}

@app.post("/api/students/preview-import")
async def preview_students_import(file: UploadFile = File(...)):
    if not file.filename or not file.filename.lower().endswith(".xlsx"): raise HTTPException(400, "Upload file Excel .xlsx")
    parsed = read_students_excel(await read_limited_upload(file)); rows = parsed["rows"]; conn = db_or_500()
    try:
        cursor = conn.cursor(); existing = set(); nis_list = [row["nis"] for row in rows]
        if nis_list:
            cursor.execute("SELECT nis FROM siswa WHERE nis = ANY(%s)", (nis_list,)); existing = {row["nis"] for row in cursor.fetchall()}
        return {"total": len(rows), "new_count": sum(row["nis"] not in existing for row in rows), "duplicate_database_count": len(existing), "duplicate_file_count": len(parsed["duplicates"]), "preview": rows[:20], "duplicates": parsed["duplicates"][:20], "skipped_rows": parsed["skipped"][:20], "sheet_summaries": parsed["sheets"]}
    finally: conn.close()

@app.post("/api/students/import")
async def import_students(file: UploadFile = File(...)):
    if not file.filename or not file.filename.lower().endswith(".xlsx"): raise HTTPException(400, "Upload file Excel .xlsx")
    parsed = read_students_excel(await read_limited_upload(file)); rows = parsed["rows"]; conn = db_or_500(); inserted = 0
    try:
        cursor = conn.cursor()
        for row in rows:
            cursor.execute("INSERT INTO siswa (nis,nisn,nama,jeniskelamin,kelas) VALUES (%s,%s,%s,%s,%s) ON CONFLICT (nis) DO NOTHING", (row["nis"], row["nisn"], row["nama"], row["jeniskelamin"], row["kelas"]))
            inserted += cursor.rowcount
        conn.commit(); return {"inserted": inserted, "skipped": len(rows) - inserted + len(parsed["duplicates"]), "skipped_rows": parsed["skipped"]}
    except psycopg2.Error:
        conn.rollback(); raise HTTPException(500, "Database error")
    finally: conn.close()

@app.get("/api/eskul/manage")
async def manage_eskul_list():
    conn = db_or_500()
    try:
        cursor = conn.cursor(); cursor.execute("SELECT e.id,e.nama_eskul,e.minimal_kelas,COUNT(s.id) AS siswa_count FROM eskul e LEFT JOIN siswa s ON s.eskul=e.id GROUP BY e.id,e.nama_eskul,e.minimal_kelas ORDER BY e.nama_eskul")
        return {"eskul": cursor.fetchall()}
    finally: conn.close()

def valid_eskul_name(name):
    name = name.strip()
    if not name or len(name) > 120: raise HTTPException(400, "Nama eskul wajib diisi dan maksimal 120 karakter")
    return name

def optional_eskul_id(value):
    if not value: return None
    if not value.isdigit(): raise HTTPException(400, "Eskul tidak valid")
    return int(value)

@app.post("/api/eskul/create")
async def create_eskul(nama_eskul: str = Form(...), minimal_kelas: int = Form(..., ge=1, le=6)):
    name = valid_eskul_name(nama_eskul); conn = db_or_500()
    try:
        cursor = conn.cursor(); cursor.execute("INSERT INTO eskul (nama_eskul,minimal_kelas) SELECT %s,%s WHERE NOT EXISTS (SELECT 1 FROM eskul WHERE LOWER(nama_eskul)=LOWER(%s)) RETURNING id,nama_eskul,minimal_kelas", (name,minimal_kelas,name)); result=cursor.fetchone()
        if not result: raise HTTPException(400,"Nama eskul sudah ada")
        conn.commit(); return {"eskul":result}
    finally: conn.close()

@app.post("/api/eskul/{eskul_id}/update")
async def update_eskul(eskul_id:int,nama_eskul:str=Form(...),minimal_kelas:int=Form(...,ge=1,le=6)):
    name=valid_eskul_name(nama_eskul); conn=db_or_500()
    try:
        cursor=conn.cursor(); cursor.execute("SELECT 1 FROM eskul WHERE LOWER(nama_eskul)=LOWER(%s) AND id<>%s",(name,eskul_id))
        if cursor.fetchone(): raise HTTPException(400,"Nama eskul sudah ada")
        cursor.execute("UPDATE eskul SET nama_eskul=%s,minimal_kelas=%s WHERE id=%s RETURNING id,nama_eskul,minimal_kelas",(name,minimal_kelas,eskul_id)); result=cursor.fetchone()
        if not result: raise HTTPException(404,"Eskul tidak ditemukan")
        cursor.execute("SELECT id,kelas FROM siswa WHERE eskul=%s FOR UPDATE",(eskul_id,)); assigned=cursor.fetchall(); invalid=[]
        for student in assigned:
            try:
                if grade_from_kelas(student["kelas"]) < minimal_kelas: invalid.append(student["id"])
            except HTTPException:
                invalid.append(student["id"])
        if invalid: cursor.execute("UPDATE siswa SET eskul=NULL WHERE id=ANY(%s)",(invalid,))
        affected=len(invalid); conn.commit(); return {"eskul":result,"affected_students":affected}
    except psycopg2.Error:
        conn.rollback(); raise HTTPException(500,"Database error")
    finally: conn.close()

@app.delete("/api/eskul/{eskul_id}")
async def delete_eskul(eskul_id:int):
    conn=db_or_500()
    try:
        cursor=conn.cursor(); cursor.execute("SELECT nama_eskul FROM eskul WHERE id=%s FOR UPDATE",(eskul_id,)); item=cursor.fetchone()
        if not item: raise HTTPException(404,"Eskul tidak ditemukan")
        cursor.execute("UPDATE siswa SET eskul=NULL WHERE eskul=%s",(eskul_id,)); affected=cursor.rowcount
        cursor.execute("DELETE FROM eskul WHERE id=%s",(eskul_id,)); conn.commit()
        return {"success":True,"name":item["nama_eskul"],"affected_students":affected}
    except psycopg2.Error:
        conn.rollback(); raise HTTPException(500,"Database error")
    finally: conn.close()

@app.get("/api/students/manage")
@app.get("/api/registrations")
async def manage_students(page:int=1,page_size:int=25,search:str="",kelas:str="",eskul_id:int|None=None,status:str=""):
    if page<1 or page_size<1 or page_size>100: raise HTTPException(400,"Pagination tidak valid")
    where,params=filters_sql(search,kelas,eskul_id,status); conn=db_or_500()
    try:
        cursor=conn.cursor(); cursor.execute("SELECT COUNT(*) total,COUNT(*) FILTER (WHERE eskul IS NOT NULL) assigned,COUNT(*) FILTER (WHERE eskul IS NULL) unassigned,COUNT(DISTINCT kelas) classes FROM siswa") ; summary=cursor.fetchone()
        cursor.execute("SELECT COUNT(*) total FROM siswa s"+where,params); total=cursor.fetchone()["total"]; pages=max(1,(total+page_size-1)//page_size)
        cursor.execute("SELECT s.id,COALESCE(s.nis,'') nis,COALESCE(s.nisn,'') nisn,COALESCE(s.nama,'') nama,COALESCE(s.jeniskelamin,'') jeniskelamin,COALESCE(s.kelas,'') kelas,s.eskul,e.nama_eskul FROM siswa s LEFT JOIN eskul e ON e.id=s.eskul"+where+" ORDER BY s.kelas,s.nama LIMIT %s OFFSET %s",params+[page_size,(page-1)*page_size]); items=cursor.fetchall()
        cursor.execute("SELECT DISTINCT kelas FROM siswa ORDER BY kelas"); classes=[row["kelas"] for row in cursor.fetchall()]
        cursor.execute("SELECT id,nama_eskul,minimal_kelas FROM eskul ORDER BY nama_eskul"); eskul=cursor.fetchall()
        return {"items":items,"students":items,"total":total,"page":page,"pages":pages,"summary":summary,"options":{"kelas":classes,"eskul":eskul}}
    finally: conn.close()

@app.get("/api/registrations/export")
async def export_registrations(search:str="",kelas:str="",eskul_id:int|None=None,status:str=""):
    where,params=filters_sql(search,kelas,eskul_id,status); conn=db_or_500()
    try:
        cursor=conn.cursor(); cursor.execute("SELECT s.nis,s.nisn,s.nama,s.jeniskelamin,s.kelas,COALESCE(e.nama_eskul,'Belum Daftar') nama_eskul FROM siswa s LEFT JOIN eskul e ON e.id=s.eskul"+where+" ORDER BY s.kelas,s.nama",params); rows=cursor.fetchall()
    finally: conn.close()
    output=io.StringIO(); writer=csv.writer(output); writer.writerow(["NIS","NISN","Nama Siswa","Jenis Kelamin","Kelas","Ekstrakurikuler"]); writer.writerows([row.values() for row in rows]); output.seek(0)
    return StreamingResponse(iter([output.getvalue()]),media_type="text/csv; charset=utf-8",headers={"Content-Disposition":"attachment; filename=registrasi_eskul.csv"})

@app.post("/api/students/create")
async def create_student(nis:str=Form(...),nisn:str=Form(...),nama:str=Form(...),jeniskelamin:str=Form(...),kelas:str=Form(...),eskul_id:str=Form("")):
    values=clean_student(nis,nisn,nama,jeniskelamin,kelas); selected_eskul=optional_eskul_id(eskul_id); conn=db_or_500()
    try:
        cursor=conn.cursor(); validate_assignment(cursor,values[4],selected_eskul); cursor.execute("INSERT INTO siswa (nis,nisn,nama,jeniskelamin,kelas,eskul) VALUES (%s,%s,%s,%s,%s,%s) RETURNING id",(*values,selected_eskul)); result=cursor.fetchone(); conn.commit(); return result
    except psycopg2.IntegrityError:
        conn.rollback(); raise HTTPException(400,"NIS sudah ada atau eskul tidak valid")
    finally: conn.close()

@app.post("/api/students/{student_id}/update")
async def update_student(student_id:int,nis:str=Form(...),nisn:str=Form(...),nama:str=Form(...),jeniskelamin:str=Form(...),kelas:str=Form(...),eskul_id:str=Form("")):
    values=clean_student(nis,nisn,nama,jeniskelamin,kelas); selected_eskul=optional_eskul_id(eskul_id); conn=db_or_500()
    try:
        cursor=conn.cursor(); validate_assignment(cursor,values[4],selected_eskul); cursor.execute("UPDATE siswa SET nis=%s,nisn=%s,nama=%s,jeniskelamin=%s,kelas=%s,eskul=%s WHERE id=%s",(*values,selected_eskul,student_id))
        if not cursor.rowcount: raise HTTPException(404,"Siswa tidak ditemukan")
        conn.commit(); return {"success":True}
    except psycopg2.IntegrityError:
        conn.rollback(); raise HTTPException(400,"NIS sudah ada atau eskul tidak valid")
    finally: conn.close()

@app.delete("/api/students/{student_id}")
async def delete_student(student_id:int):
    conn=db_or_500()
    try:
        cursor=conn.cursor(); cursor.execute("DELETE FROM siswa WHERE id=%s",(student_id,))
        if not cursor.rowcount: raise HTTPException(404,"Siswa tidak ditemukan")
        conn.commit(); return {"success":True}
    finally: conn.close()

@app.post("/api/students/bulk-delete")
async def bulk_delete_students(ids:List[int]=Body(...)):
    if not ids or len(ids)>100: raise HTTPException(400,"Pilih 1 sampai 100 siswa")
    conn=db_or_500()
    try:
        cursor=conn.cursor(); cursor.execute("DELETE FROM siswa WHERE id=ANY(%s)",(list(set(ids)),)); deleted=cursor.rowcount; conn.commit(); return {"deleted":deleted}
    finally: conn.close()

@app.post("/api/submit")
async def submit_form(siswa_id:int=Form(...),eskul_id:int=Form(...)):
    conn=db_or_500()
    try:
        cursor=conn.cursor(); cursor.execute("SELECT nama_eskul,minimal_kelas FROM eskul WHERE id=%s FOR SHARE",(eskul_id,)); eskul=cursor.fetchone()
        if not eskul: raise HTTPException(404,"Siswa atau eskul tidak ditemukan")
        cursor.execute("SELECT kelas,nama FROM siswa WHERE id=%s FOR UPDATE",(siswa_id,)); student=cursor.fetchone()
        if not student: raise HTTPException(404,"Siswa atau eskul tidak ditemukan")
        if grade_from_kelas(student["kelas"]) < eskul["minimal_kelas"]: raise HTTPException(400,f"{eskul['nama_eskul']} hanya untuk kelas {eskul['minimal_kelas']}+")
        cursor.execute("UPDATE siswa SET eskul=%s WHERE id=%s AND eskul IS NULL",(eskul_id,siswa_id))
        if not cursor.rowcount: raise HTTPException(409,"Siswa sudah memilih ekstrakurikuler. Pilihan tidak diubah.")
        conn.commit(); return {"success":True,"message":f"Berhasil mendaftarkan {student['nama']} ke eskul {eskul['nama_eskul']}","data":{"nama":student["nama"],"nama_eskul":eskul["nama_eskul"],"kelas":student["kelas"]}}
    except psycopg2.Error:
        conn.rollback(); raise HTTPException(500,"Database error")
    finally: conn.close()

@app.get("/registrations",response_class=HTMLResponse)
async def view_registrations(request:Request):
    return templates.TemplateResponse("registrations.html",{"request":request,"login":False})

if __name__=="__main__":
    import uvicorn
    uvicorn.run(app,host="0.0.0.0",port=8000)
